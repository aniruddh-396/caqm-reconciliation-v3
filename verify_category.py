"""
Independent verifier for Module 2 (category-wise).

Re-derives expected values straight from sectoral_analysis using a plain dict
keyed by (state, category), without importing reconcile_category.py, then
asserts the corrected output matches cell by cell.
"""
import sys
import re
import openpyxl

EDIT, BASE, OUT = sys.argv[1:4]
ERR = {"#DIV/0!", "#N/A", "#VALUE!", "#REF!", "#NAME?", "N/A"}


def norm(t):
    if t is None:
        return ""
    s = re.sub(r"\s+", " ", str(t).replace("\u00a0", " ")).strip().casefold()
    return re.sub(r"[.\-_,'\"]+", "", s)


def num(v):
    if v is None or (isinstance(v, str) and v.strip().upper() in ERR):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


# ---- build expected values from the base ----------------------------------
bws = openpyxl.load_workbook(BASE, data_only=True).worksheets[0]
exp = {}
state = ""
for r in range(3, bws.max_row + 1):
    st = bws.cell(r, 1).value
    cat = bws.cell(r, 2).value
    if st and not cat and norm(st).endswith("total"):
        continue
    if st and not norm(st).endswith("total"):
        state = str(st).strip()
    if not cat:
        continue
    exp[(norm(state), norm(cat))] = (
        int(num(bws.cell(r, 3).value) or 0),   # pending
        int(num(bws.cell(r, 4).value) or 0),   # resolved
        int(num(bws.cell(r, 5).value) or 0),   # grand total
    )

# ---- resolve merged state cells in the output -----------------------------
out = openpyxl.load_workbook(OUT, data_only=True).worksheets[0]
merges = openpyxl.load_workbook(OUT).worksheets[0].merged_cells.ranges


def state_at(r):
    v = out.cell(r, 2).value
    if v:
        return str(v).strip()
    for rng in merges:
        if rng.min_col <= 2 <= rng.max_col and rng.min_row <= r <= rng.max_row:
            a = out.cell(rng.min_row, rng.min_col).value
            return str(a).strip() if a else ""
    return ""


fails, checked, used = [], 0, set()
for r in range(3, out.max_row + 1):
    cat = out.cell(r, 3).value
    if not cat:
        continue
    st = state_at(r)
    key = (norm(st), norm(cat))
    if key not in exp:
        fails.append(f"row {r}: ({st} / {cat}) has no base counterpart")
        continue
    used.add(key)
    p, res_, g = exp[key]
    got_recv = num(out.cell(r, 4).value)
    got_pend = num(out.cell(r, 5).value)
    got_res = num(out.cell(r, 6).value)
    checked += 3
    if got_recv != g:
        fails.append(f"row {r} {st}/{cat}: Received expected {g}, got {got_recv}")
    if got_pend != p:
        fails.append(f"row {r} {st}/{cat}: Pending expected {p}, got {got_pend}")
    if got_res != res_:
        fails.append(f"row {r} {st}/{cat}: Resolved expected {res_}, got {got_res}")

missing = set(exp) - used
if missing:
    fails.append(f"{len(missing)} base categories never matched in output")

# ---- total rows must sum their member category rows ------------------------
blocks = {}
cur_first = None
for r in range(3, out.max_row + 1):
    st_raw = out.cell(r, 2).value
    cat = out.cell(r, 3).value
    if st_raw and norm(st_raw).endswith("total") and not cat:
        blocks[r] = (cur_first, r - 1)
        cur_first = None
    elif cat:
        if cur_first is None:
            cur_first = r

for trow, (a, b) in blocks.items():
    if a is None:
        continue
    for col, nm in ((4, "Received"), (5, "Pending"), (6, "Resolved")):
        checked += 1
        want = sum(num(out.cell(x, col).value) or 0 for x in range(a, b + 1))
        got = num(out.cell(trow, col).value)
        if got != want:
            fails.append(f"TOTAL row {trow} {nm}: expected {want}, got {got}")

# grand total (last total-looking row) vs sum of state totals
gt_row = max((r for r in range(3, out.max_row + 1)
             if (v := out.cell(r, 2).value) and norm(v) == "grand total"), default=None)
if gt_row:
    state_totals = [r for r in blocks if r != gt_row]
    for col, nm in ((4, "Received"), (5, "Pending"), (6, "Resolved")):
        checked += 1
        want = sum(num(out.cell(x, col).value) or 0 for x in state_totals)
        got = num(out.cell(gt_row, col).value)
        if got != want:
            fails.append(f"GRAND TOTAL row {gt_row} {nm}: expected {want}, got {got}")

print(f"assertions checked: {checked}")
if fails:
    print(f"FAILURES: {len(fails)}")
    for f in fails:
        print("   x", f)
    sys.exit(1)
print("ALL CHECKS PASSED")
