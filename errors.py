"""
Shared error reporting and format validation for all CAQM modules.

Goal: when something is wrong, the message must say WHICH file, WHICH sheet,
WHICH row/column, what was EXPECTED and what was actually FOUND - never a bare
"KeyError" or a silently blank cell.

Two kinds of problem are reported:

  FormatError  - the workbook's layout is not what the module expects
                 (missing/renamed column, header row not found, sheet empty).
                 Reported per offending row/column.

  DataError    - the layout is fine but the contents cannot be reconciled
                 (an agency in one file with no counterpart in the other,
                 a duplicate key, a non-numeric value in a count column).

Both carry a list of `Issue` records so the UI can show them as a table and the
CLI can print them as lines.
"""

from __future__ import annotations

import os
import re
from dataclasses import dataclass, field, asdict

import openpyxl


# ---------------------------------------------------------------------------
# Issue record
# ---------------------------------------------------------------------------

@dataclass
class Issue:
    file: str                 # file name the problem is in
    problem: str              # short category, e.g. "Missing column"
    detail: str               # human sentence explaining it
    sheet: str = ""
    row: object = ""          # row number, or "" when not row-specific
    column: str = ""          # column letter/name, or ""
    expected: str = ""
    found: str = ""
    fix: str = ""             # what the user should do

    def as_dict(self):
        return asdict(self)

    def line(self):
        loc = []
        if self.sheet:
            loc.append(f"sheet '{self.sheet}'")
        if self.row != "" and self.row is not None:
            loc.append(f"row {self.row}")
        if self.column:
            loc.append(f"column {self.column}")
        where = ", ".join(loc)
        s = f"[{self.file}" + (f" / {where}" if where else "") + f"] {self.problem}: {self.detail}"
        if self.expected or self.found:
            s += f"\n      expected: {self.expected or '-'}\n      found   : {self.found or '-'}"
        if self.fix:
            s += f"\n      fix     : {self.fix}"
        return s


class CaqmError(Exception):
    """Base for all reportable problems. Carries structured issues."""

    def __init__(self, issues, headline=""):
        self.issues = issues if isinstance(issues, list) else [issues]
        self.headline = headline or self.default_headline
        super().__init__(self.render())

    default_headline = "The run was stopped"

    def render(self):
        body = "\n".join("  - " + i.line() for i in self.issues)
        return f"{self.headline}\n{body}"

    def table(self):
        return [i.as_dict() for i in self.issues]


class FormatError(CaqmError):
    default_headline = ("The file's format is not what this module expects, so "
                        "nothing was changed")


class DataError(CaqmError):
    default_headline = ("The data could not be reconciled, so nothing was "
                        "changed (no value is ever guessed)")


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def norm(text) -> str:
    if text is None:
        return ""
    s = re.sub(r"[\s\n\r\t]+", " ", str(text).replace("\u00a0", " ")).strip().casefold()
    s = re.sub(r"[.\-_,'\"]+", "", s)
    return re.sub(r"\s+", " ", s)


def col_letter(idx: int) -> str:
    return openpyxl.utils.get_column_letter(idx)


def _row_preview(ws, row, limit=12):
    vals = []
    for c in range(1, min(limit, ws.max_column) + 1):
        v = ws.cell(row, c).value
        if v is not None:
            vals.append(f"{col_letter(c)}={str(v)[:28]}")
    return ", ".join(vals) if vals else "(empty row)"


# ---------------------------------------------------------------------------
# Format validation
# ---------------------------------------------------------------------------

def open_sheet(path, label=None):
    """Open the first worksheet, with a clear error if the file is unusable."""
    name = label or os.path.basename(path)
    if not os.path.exists(path):
        raise FormatError(Issue(file=name, problem="File not found",
                                detail=f"No file at {path}",
                                fix="Upload the file again."))
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        raise FormatError(Issue(
            file=name, problem="File could not be opened",
            detail=str(e)[:200],
            fix="Make sure it is a real .xlsx file (not .xls, not CSV renamed) "
                "and is not password protected or currently open in Excel."))
    ws = wb.worksheets[0]
    if ws.max_row < 2:
        raise FormatError(Issue(file=name, sheet=ws.title, problem="Sheet is empty",
                                detail=f"Only {ws.max_row} row(s) found.",
                                fix="Check you uploaded the right file."))
    return ws


def find_header(path, required, label=None, search_rows=15):
    """
    Locate the header row and map required columns to column numbers.

    `required` is a list of (key, [acceptable header names...]).
    Every key must be found on one single row, or a FormatError is raised that
    names the row it got closest on and exactly which headers were missing.

    Returns (worksheet, header_row_number, {key: column_number}).
    """
    name = label or os.path.basename(path)
    ws = open_sheet(path, name)

    best_row, best_hits, best_map = None, -1, {}
    for r in range(1, min(search_rows, ws.max_row) + 1):
        present = {}
        for c in range(1, ws.max_column + 1):
            h = norm(ws.cell(r, c).value)
            if h and h not in present:
                present[h] = c

        found, hits = {}, 0
        for key, names in required:
            col = None
            for n in names:                       # exact match first
                if norm(n) in present:
                    col = present[norm(n)]
                    break
            if col is None:                       # then contains match
                for n in names:
                    nn = norm(n)
                    for h, c in present.items():
                        if nn and (nn in h or h in nn):
                            col = c
                            break
                    if col:
                        break
            if col:
                found[key] = col
                hits += 1
        if hits > best_hits:
            best_row, best_hits, best_map = r, hits, found
        if hits == len(required):
            return ws, r, found

    # nothing matched fully - report precisely what is missing
    missing = [key for key, _ in required if key not in best_map]
    issues = []
    for key in missing:
        names = dict(required)[key]
        issues.append(Issue(
            file=name, sheet=ws.title, row=best_row,
            problem="Missing column",
            detail=f"The column '{names[0]}' is not present in the header row.",
            expected=" or ".join(names),
            found=_row_preview(ws, best_row) if best_row else "-",
            fix=f"Add a '{names[0]}' column header, or rename the existing "
                f"column to '{names[0]}'. Column order does not matter."))
    raise FormatError(issues, headline=(
        f"'{name}' does not have the expected columns, so nothing was changed. "
        f"The closest match was row {best_row} "
        f"({best_hits} of {len(required)} columns found)."))


def check_numeric(ws, row, col, file_name, field_label, allow_blank=True,
                  allow_text=()):
    """Validate one cell is numeric; return the number or raise a DataError
    that names the exact cell."""
    v = ws.cell(row, col).value
    if v is None and allow_blank:
        return 0.0
    if isinstance(v, str):
        t = v.strip()
        if t == "" and allow_blank:
            return 0.0
        if norm(t) in {norm(a) for a in allow_text}:
            return None
        try:
            return float(t.replace(",", ""))
        except ValueError:
            raise DataError(Issue(
                file=file_name, sheet=ws.title, row=row, column=col_letter(col),
                problem="Value is not a number",
                detail=f"{field_label} contains text that cannot be read as a number.",
                expected="a number", found=repr(v)[:60],
                fix="Correct this cell in the source file, then run again."))
    if isinstance(v, (int, float)):
        return float(v)
    raise DataError(Issue(
        file=file_name, sheet=ws.title, row=row, column=col_letter(col),
        problem="Unexpected value type",
        detail=f"{field_label} holds a value of type {type(v).__name__}.",
        expected="a number", found=repr(v)[:60],
        fix="Correct this cell in the source file, then run again."))


def duplicate_issue(file_name, sheet, row, first_row, key_label, key_value):
    return Issue(
        file=file_name, sheet=sheet, row=row,
        problem="Duplicate entry",
        detail=f"{key_label} '{key_value}' already appeared on row {first_row}.",
        expected="each entry once",
        found=f"rows {first_row} and {row}",
        fix="Remove or merge the duplicate row in the source file.")


def unmatched_issue(file_name, sheet, row, label, value, other_file, hint=""):
    return Issue(
        file=file_name, sheet=sheet, row=row,
        problem="No matching entry",
        detail=f"{label} '{value}' has no counterpart in {other_file}.",
        expected=f"a matching row in {other_file}",
        found="no match",
        fix=hint or ("Check the spelling in both files, or add the missing row. "
                     "Matching ignores case, spacing and punctuation."))
