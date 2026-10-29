"""
Independent verifier.

Deliberately does NOT import reconcile.py. It re-derives every expected value
straight from the base workbooks using plain dictionaries, then asserts the
corrected output matches. If reconcile.py has a bug, this should disagree.
"""
import sys
from collections import defaultdict
import openpyxl
from mapping import HANDLE_TO_AGENCY, EXPECTED_MISSING, PERIOD_TOLERANCE_DAYS, RECEIVED_MODE

EDIT, COUNTS, PERIODS, OUT = sys.argv[1:5]

ERR = {"#DIV/0!", "#N/A", "#VALUE!", "#REF!", "#NAME?", "N/A"}


def num(v):
    if v is None or (isinstance(v, str) and v.strip().upper() in ERR):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


# ---- expected values, built by hand from the base files --------------------
exp_cnt = defaultdict(lambda: [0, 0, 0])      # agency -> [pending, resolved, grand]
exp_res = defaultdict(list)
exp_pen = defaultdict(list)

ws = openpyxl.load_workbook(COUNTS, data_only=True).worksheets[0]
for r in range(3, ws.max_row + 1):
    h = ws.cell(r, 2).value
    if not h:
        continue
    a = HANDLE_TO_AGENCY[str(h).strip()]
    exp_cnt[a][0] += num(ws.cell(r, 3).value) or 0
    exp_cnt[a][1] += num(ws.cell(r, 4).value) or 0
    exp_cnt[a][2] += num(ws.cell(r, 5).value) or 0

ws = openpyxl.load_workbook(PERIODS, data_only=True).worksheets[0]
for r in range(2, ws.max_row + 1):
    h = ws.cell(r, 2).value
    if not h:
        continue
    a = HANDLE_TO_AGENCY[str(h).strip()]
    if (v := num(ws.cell(r, 3).value)) is not None:
        exp_res[a].append(v)
    if (v := num(ws.cell(r, 4).value)) is not None:
        exp_pen[a].append(v)


def mean_round(vals):
    if not vals:
        return None
    from decimal import Decimal, ROUND_HALF_UP
    return int(Decimal(str(sum(vals) / len(vals))).quantize(Decimal(1), rounding=ROUND_HALF_UP))


# ---- check the produced workbook -------------------------------------------
out = openpyxl.load_workbook(OUT, data_only=True).worksheets[0]
fml = openpyxl.load_workbook(OUT).worksheets[0]

fails, checked = [], 0
for r in range(3, out.max_row + 1):
    agency = out.cell(r, 3).value
    if not agency:
        continue
    agency = str(agency).strip()
    if agency in EXPECTED_MISSING:
        continue
    if agency not in exp_cnt and agency not in exp_res:
        fails.append(f"row {r}: '{agency}' has no base counterpart")
        continue

    p, res_, g = exp_cnt[agency]
    got_p, got_r, got_d = num(out.cell(r, 5).value), num(out.cell(r, 7).value), num(out.cell(r, 4).value)

    checked += 3
    if got_p != p:
        fails.append(f"row {r} {agency}: Pending expected {p}, got {got_p}")
    if got_r != res_:
        fails.append(f"row {r} {agency}: Resolved expected {res_}, got {got_r}")

    want_d = p + res_ if RECEIVED_MODE == "formula" else g
    if got_d != want_d:
        fails.append(f"row {r} {agency}: Received expected {want_d}, got {got_d}")

    for col, vals, name in ((10, exp_res[agency], "AvgResolution"),
                            (11, exp_pen[agency], "AvgPendency")):
        checked += 1
        want = mean_round(vals)
        got_raw = out.cell(r, col).value
        got = num(got_raw)
        if want is None:
            if got is not None:
                fails.append(f"row {r} {agency}: {name} expected N/A, got {got_raw!r}")
        elif got is None:
            fails.append(f"row {r} {agency}: {name} expected {want}, got {got_raw!r}")
        elif abs(got - want) >= PERIOD_TOLERANCE_DAYS:
            fails.append(f"row {r} {agency}: {name} expected {want}, got {got}")

    # percentage columns must stay consistent with the counts
    for pc, numer, nm in ((6, p, "pending%"), (8, res_, "resolved%")):
        v = out.cell(r, pc).value
        if isinstance(v, str) and v.startswith("("):
            checked += 1
            # Excel ROUND() is half-away-from-zero; Python round() is half-to-even
            from decimal import Decimal, ROUND_HALF_UP
            want_pct = (int(Decimal(str(numer / want_d * 100)).quantize(
                Decimal(1), rounding=ROUND_HALF_UP)) if want_d else None)
            if want_pct is not None and v != f"({want_pct}%)":
                fails.append(f"row {r} {agency}: {nm} expected ({want_pct}%), got {v}")

# total rows must equal the sum of their member rows
for trow, rng in ((19, range(3, 19)), (32, range(20, 32)), (34, range(33, 34)),
                  (47, range(35, 47))):
    for col, nm in ((5, "Pending"), (7, "Resolved"), (4, "Received")):
        checked += 1
        want = sum(num(out.cell(x, col).value) or 0 for x in rng)
        got = num(out.cell(trow, col).value)
        if got != want:
            fails.append(f"TOTAL row {trow} {nm}: expected {want}, got {got}")

# grand total (row 50) = state totals + loose agency rows 48, 49
for col, nm in ((5, "Pending"), (7, "Resolved"), (4, "Received")):
    checked += 1
    want = sum(num(out.cell(x, col).value) or 0 for x in (19, 32, 34, 47, 48, 49))
    got = num(out.cell(50, col).value)
    if got != want:
        fails.append(f"GRAND TOTAL row 50 {nm}: expected {want}, got {got}")

print(f"assertions checked: {checked}")
if fails:
    print(f"FAILURES: {len(fails)}")
    for f in fails:
        print("   x", f)
    sys.exit(1)
print("ALL CHECKS PASSED")
