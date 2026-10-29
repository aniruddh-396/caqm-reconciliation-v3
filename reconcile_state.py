"""
Module 3 - Social Media State-wise generator
=============================================

Splits the parent workbook (2025.xlsx) into one workbook per state by pouring
fresh numbers into fixed per-state templates.

    Delhi      <- parent agencies of Delhi (3 MCD zones summed into one row)
    UP         <- parent Uttar Pradesh agencies
    Haryana    <- parent Haryana agencies
    Rajasthan  <- parent Rajasthan agency

How it works
------------
* The four template files in ./templates define the exact layout of each state
  file - the agency list, the grouping, where Sub Total / Total / Grand Total
  rows sit, and all formatting. Those templates never change; only the numbers
  and the end date are refreshed on every run.
* For each agency row in a template, the agency name is matched to the parent
  (whitespace/punctuation-insensitive) and the counts are copied in:
      C = Total Received, D = Resolved, F = Not Resolved,
      E = (% Resolved),   G = (% Not Resolved)   [recomputed]
* "Municipal Corporation of Delhi" is the one aggregate row: East + North +
  South Delhi Municipal Corporation are summed (see AGGREGATES).
* Sub Total / Total rows are recomputed as the sum of their block; Grand Total
  as the sum of every agency row in the sheet.
* An unmatched template agency ABORTS the run (never silently blank).
* NHAI and Northern Railway exist in the parent but appear in no template, so
  they are simply never pulled - exactly as required.

Nothing is invented: if the parent has a value it is copied; percentages are
the only derived figures and they follow round-half-up, matching the parent.
"""

from __future__ import annotations

import os
import re
import shutil
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP

import openpyxl

from reconcile import norm, as_number, is_error_value
from errors import Issue, FormatError, DataError, find_header, duplicate_issue
from datefmt import heading_date, file_date, replace_end_date_in_title, ordinal


# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

TEMPLATE_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "templates")

STATES = [
    {"key": "delhi",     "prefix": "Delhi",     "template": "delhi.xlsx"},
    {"key": "up",        "prefix": "UP",        "template": "up.xlsx"},
    {"key": "haryana",   "prefix": "Haryana",   "template": "haryana.xlsx"},
    {"key": "rajasthan", "prefix": "Rajasthan", "template": "rajasthan.xlsx"},
]

# Output agency name -> the parent agency names it aggregates (summed).
AGGREGATES = {
    norm("Municipal Corporation of Delhi"): [
        "East Delhi Municipal Corporation",
        "North Delhi Municipal Corporation",
        "South Delhi Municipal Corporation",
    ],
}

FIXED_START_FILE = "01-01-2025"          # start date is fixed
COLS = {"received": 3, "resolved": 4, "res_pct": 5, "notres": 6, "notres_pct": 7}


# ---------------------------------------------------------------------------
# Date helpers
# ---------------------------------------------------------------------------

# ---------------------------------------------------------------------------
# Parent reader
# ---------------------------------------------------------------------------

def read_parent(path):
    """norm(agency) -> {'name','received','resolved','notres','row'}."""
    fname = os.path.basename(path)
    spec = [("sno", ["S. No.", "S.No.", "Sr. No.", "Sr No"]),
            ("agency", ["Agency"])]
    ws, hrow, _ = find_header(path, spec, fname)

    out = {}
    for r in range(hrow + 1, ws.max_row + 1):
        name = ws.cell(r, 2).value
        if not name or norm(name).endswith("total"):
            continue
        key = norm(name)
        rec = {
            "name": re.sub(r"\s+", " ", str(name).replace("\n", " ")).strip(),
            "received": int(as_number(ws.cell(r, COLS["received"]).value) or 0),
            "resolved": int(as_number(ws.cell(r, COLS["resolved"]).value) or 0),
            "notres": int(as_number(ws.cell(r, COLS["notres"]).value) or 0),
            "row": r,
        }
        if key in out:
            raise DataError(duplicate_issue(fname, ws.title, r, out[key]["row"],
                                            "Agency", rec["name"]))
        out[key] = rec
    return out


# ---------------------------------------------------------------------------
# Percentage
# ---------------------------------------------------------------------------

def pct(part: int, whole: int):
    """'(NN%)' rounded half-up, or 0 when the denominator is 0 (matches the
    state templates, which show 0 rather than #DIV/0! for empty agencies)."""
    if whole == 0:
        return 0
    p = int(Decimal(str(part / whole * 100)).quantize(Decimal(1), rounding=ROUND_HALF_UP))
    return f"({p}%)"


# ---------------------------------------------------------------------------
# Generate one state file
# ---------------------------------------------------------------------------

def generate_state(state, parent, end_dt, outdir):
    template = os.path.join(TEMPLATE_DIR, state["template"])
    if not os.path.exists(template):
        raise FormatError(Issue(
            file=f"templates/{state['template']}", problem="Template missing",
            detail=f"The fixed layout for {state['prefix']} is not in templates/.",
            fix="Make sure the whole templates/ folder sits next to app.py."))

    prefix = state["prefix"]
    fname = f"{prefix}_Social Media Agency-wise Complaint Redressal " \
            f"({FIXED_START_FILE} - {file_date(end_dt)}).xlsx"
    out_path = os.path.join(outdir, fname)
    shutil.copyfile(template, out_path)

    wb = openpyxl.load_workbook(out_path)
    ws = wb.worksheets[0]

    # locate header row (has 'S. No.' + 'Agency')
    hrow = None
    for r in range(1, 12):
        cells = [norm(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        if any(c == "agency" for c in cells) and any("no" in c and c for c in cells):
            hrow = r
            break
    if hrow is None:
        raise FormatError(Issue(
            file=f"templates/{state['template']}", problem="Header row not found",
            detail="Could not find a row containing both 'S. No.' and 'Agency'.",
            expected="a header row with 'S. No.' and 'Agency'",
            fix="Restore the original template file."))

    # ---- title date ----
    a1 = ws.cell(1, 1).value
    if isinstance(a1, str):
        ws.cell(1, 1).value = replace_end_date_in_title(a1, end_dt)

    changes, unmatched = [], []
    all_agency_rows, block_rows = [], []

    def write_counts(r, received, resolved, notres, label, is_total=False):
        ws.cell(r, COLS["received"]).value = int(received)
        ws.cell(r, COLS["resolved"]).value = int(resolved)
        ws.cell(r, COLS["notres"]).value = int(notres)
        ws.cell(r, COLS["res_pct"]).value = pct(resolved, received)
        ws.cell(r, COLS["notres_pct"]).value = pct(notres, received)
        changes.append({"row": r, "agency": label, "received": int(received),
                        "resolved": int(resolved), "notres": int(notres),
                        "kind": "TOTAL" if is_total else "agency"})

    for r in range(hrow + 1, ws.max_row + 1):
        name = ws.cell(r, 2).value
        if not name:
            continue
        n = norm(name)

        # total-type row
        if n.endswith("total"):
            if n == "grand total":
                rows = all_agency_rows
            else:
                rows = block_rows
            recv = sum(ws.cell(x, COLS["received"]).value or 0 for x in rows)
            res = sum(ws.cell(x, COLS["resolved"]).value or 0 for x in rows)
            nres = sum(ws.cell(x, COLS["notres"]).value or 0 for x in rows)
            write_counts(r, recv, res, nres,
                         re.sub(r"\s+", " ", str(name)).strip(), is_total=True)
            block_rows = []
            continue

        # agency row
        clean = re.sub(r"\s+", " ", str(name).replace("\n", " ")).strip()
        if n in AGGREGATES:
            parts = AGGREGATES[n]
            missing = [p for p in parts if norm(p) not in parent]
            if missing:
                unmatched.append({"row": r, "agency": clean,
                                  "status": f"aggregate parts missing from parent: {missing}"})
                continue
            recv = sum(parent[norm(p)]["received"] for p in parts)
            res = sum(parent[norm(p)]["resolved"] for p in parts)
            nres = sum(parent[norm(p)]["notres"] for p in parts)
        else:
            rec = parent.get(n)
            if rec is None:
                unmatched.append({"row": r, "agency": clean,
                                  "status": "NOT FOUND in parent"})
                continue
            recv, res, nres = rec["received"], rec["resolved"], rec["notres"]

        write_counts(r, recv, res, nres, clean)
        all_agency_rows.append(r)
        block_rows.append(r)

    if unmatched:
        raise DataError([Issue(
            file=f"templates/{state['template']}", sheet=ws.title, row=u["row"],
            column="Agency", problem="No matching entry",
            detail=f"'{u['agency']}' ({u['status']}) could not be found in the "
                   f"parent workbook, so its row would have been left blank.",
            expected="a row with this agency name in the parent workbook",
            found=u["status"],
            fix="Check the agency name matches the parent exactly (case, spacing "
                "and punctuation are ignored), or add it to the parent.")
            for u in unmatched],
            headline=f"{prefix} could not be generated - some agencies in the "
                     f"fixed layout are not present in the parent workbook")

    wb.save(out_path)
    return {"state": prefix, "output": out_path, "filename": fname,
            "changes": changes}


# ---------------------------------------------------------------------------
# Generate all states
# ---------------------------------------------------------------------------

def generate_all(parent_path, end_dt, outdir):
    os.makedirs(outdir, exist_ok=True)
    parent = read_parent(parent_path)
    results = [generate_state(s, parent, end_dt, outdir) for s in STATES]
    return {"parent": parent, "results": results}


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main():
    import argparse
    p = argparse.ArgumentParser(description="Generate CAQM state-wise workbooks from the parent.")
    p.add_argument("--parent", required=True, help="parent workbook (e.g. 2025.xlsx)")
    p.add_argument("--end-date", required=True, help="end date YYYY-MM-DD")
    p.add_argument("--outdir", default="./out_states")
    a = p.parse_args()
    end_dt = datetime.strptime(a.end_date, "%Y-%m-%d").date()
    res = generate_all(a.parent, end_dt, a.outdir)
    for r in res["results"]:
        n_agency = sum(1 for c in r["changes"] if c["kind"] == "agency")
        print(f"{r['state']:10} -> {r['filename']}   ({n_agency} agencies)")


if __name__ == "__main__":
    main()
